"""Converte docs/banco-di-prova-modelli.xlsx in campioni.json.

    uv run python scripts/campioni_da_excel.py

Chi compila la verità dei 10 capi campione lavora nel foglio Excel, mai nel
JSON: questo script legge la scheda «Verità capi campione» e riscrive
tests/fixtures/campioni/campioni.json, validando ogni riga contro
`CampioneValutazione` — un errore di battitura si scopre qui, non a metà di
un run a pagamento con `scripts/valuta_modelli.py`.

Convenzione di ogni colonna attributo (tipo, colore, materiale, fantasia,
stagione, vestibilita, lavaggio), spiegata anche nel foglio «Leggimi»:
  - vuoto            -> non valutato (il campo non entra nel punteggio)
  - "ASSENTE"        -> il modello deve tacere: è la risposta giusta
  - "valore"         -> il valore atteso
  - "valore/sinonimo" -> il valore atteso, più sinonimi accettati come «vicino»
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

RADICE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RADICE / "src"))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.worksheet.worksheet import Worksheet  # noqa: E402

from domain.models import AttributoCapo, CampioneValutazione, VeritaAttributo  # noqa: E402
from domain.valutazione import campioni_da_json  # noqa: E402

FOGLIO_DEFAULT = RADICE.parent.parent / "docs" / "banco-di-prova-modelli.xlsx"
DESTINAZIONE_DEFAULT = RADICE / "tests" / "fixtures" / "campioni" / "campioni.json"
NOME_SCHEDA = "Verità capi campione"

COLONNE_ATTRIBUTO: dict[AttributoCapo, str] = {
    AttributoCapo.TIPO: "tipo",
    AttributoCapo.COLORE: "colore",
    AttributoCapo.MATERIALE: "materiale",
    AttributoCapo.FANTASIA: "fantasia",
    AttributoCapo.STAGIONE: "stagione",
    AttributoCapo.VESTIBILITA: "vestibilita",
    AttributoCapo.LAVAGGIO: "lavaggio",
}


def _intestazioni(foglio: Worksheet) -> dict[str, int]:
    prima_riga = next(foglio.iter_rows(min_row=1, max_row=1))
    return {
        str(cella.value).strip(): indice for indice, cella in enumerate(prima_riga) if cella.value
    }


def _testo(riga: tuple[Any, ...], indici: dict[str, int], nome: str) -> str | None:
    indice = indici.get(nome)
    if indice is None:
        return None
    valore = riga[indice].value
    if valore is None:
        return None
    testo = str(valore).strip()
    return testo or None


def _verita_attributo(
    riga: tuple[Any, ...],
    indici: dict[str, int],
    nome_colonna: str,
    *,
    hex_colore: str | None,
    tolleranza: int | None,
) -> VeritaAttributo | None:
    grezzo = _testo(riga, indici, nome_colonna)
    if grezzo is None:
        return None
    if grezzo.upper() == "ASSENTE":
        return VeritaAttributo(assente=True)

    pezzi = [p.strip() for p in grezzo.split("/") if p.strip()]
    campi: dict[str, object] = {"attesi": pezzi[:1], "vicini": pezzi[1:]}
    if hex_colore:
        campi["hex"] = hex_colore
    if tolleranza is not None:
        campi["tolleranza_hex"] = tolleranza
    return VeritaAttributo.model_validate(campi)


def converti(percorso_excel: Path) -> list[CampioneValutazione]:
    cartella = load_workbook(percorso_excel, data_only=True)
    if NOME_SCHEDA not in cartella.sheetnames:
        raise SystemExit(f"«{NOME_SCHEDA}» non è una scheda di {percorso_excel}")
    foglio = cartella[NOME_SCHEDA]
    indici = _intestazioni(foglio)

    grezzo: list[object] = []
    for riga in foglio.iter_rows(min_row=2):
        id_campione = _testo(riga, indici, "id")
        if id_campione is None:
            continue

        hex_colore = _testo(riga, indici, "colore_hex")
        tolleranza_grezza = _testo(riga, indici, "colore_tolleranza")
        tolleranza = int(tolleranza_grezza) if tolleranza_grezza else None

        verita: dict[str, object] = {}
        for attributo, nome_colonna in COLONNE_ATTRIBUTO.items():
            colore = attributo is AttributoCapo.COLORE
            valore = _verita_attributo(
                riga,
                indici,
                nome_colonna,
                hex_colore=hex_colore if colore else None,
                tolleranza=tolleranza if colore else None,
            )
            if valore is not None:
                verita[attributo.value] = valore.model_dump(mode="json")

        grezzo.append(
            {
                "id": id_campione,
                "descrizione": _testo(riga, indici, "caratteristica (non modificare)") or "",
                "foto": _testo(riga, indici, "foto (non modificare)") or f"{id_campione}.jpg",
                "verita": verita,
            }
        )

    return campioni_da_json(grezzo)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, default=FOGLIO_DEFAULT)
    parser.add_argument("--destinazione", type=Path, default=DESTINAZIONE_DEFAULT)
    args = parser.parse_args()

    campioni = converti(args.excel)
    contenuto = json.dumps(
        [c.model_dump(mode="json") for c in campioni], indent=2, ensure_ascii=False
    )
    args.destinazione.write_text(contenuto + "\n", encoding="utf-8")
    print(f"{len(campioni)} campioni scritti in {args.destinazione}")


if __name__ == "__main__":
    main()
